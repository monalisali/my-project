# -*- coding: utf-8 -*-
"""Assemble Chinese chapter digests (markdown) into a styled Word document."""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

DIGESTS = [
    r"C:\D\AIWork\my-project\.work\digests\digest_A.md",
    r"C:\D\AIWork\my-project\.work\digests\digest_B.md",
    r"C:\D\AIWork\my-project\.work\digests\digest_C.md",
    r"C:\D\AIWork\my-project\.work\digests\digest_D.md",
    r"C:\D\AIWork\my-project\.work\digests\digest_E.md",
]
OUT = r"C:\D\AIWork\my-project\translation\金钱心理学_中文章节导读.docx"

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def set_font(run, name=BODY_FONT, size=12, bold=False, color=None):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    r = run._element.rPr.rFonts
    r.set(qn("w:eastAsia"), name)


def add_para(doc, text, size=12, bold=False, indent=True, space_after=6,
             line=1.4, color=None, font=BODY_FONT, align=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Pt(size * 2)
    if align is not None:
        pf.alignment = align
    run = p.add_run(text)
    set_font(run, font, size, bold, color)
    return p


def add_bullet(doc, text, size=12):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.35
    run = p.add_run(text)
    set_font(run, BODY_FONT, size)
    return p


def add_heading(doc, text, level=1):
    sizes = {1: 17, 2: 13.5}
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(16 if level == 1 else 10)
    pf.space_after = Pt(8 if level == 1 else 5)
    if level == 1:
        pf.page_break_before = True
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_font(run, HEAD_FONT, sizes[level], bold=True, color=ACCENT)
    # bottom border for h1
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        pbdr = pPr.makeelement(qn("w:pBdr"), {})
        bottom = pPr.makeelement(qn("w:bottom"), {
            qn("w:val"): "single", qn("w:sz"): "12",
            qn("w:space"): "2", qn("w:color"): "1F4E79"})
        pbdr.append(bottom)
        pPr.append(pbdr)
    return p


def main():
    doc = Document()
    # page margins
    for sec in doc.sections:
        sec.top_margin = Cm(2.4)
        sec.bottom_margin = Cm(2.4)
        sec.left_margin = Cm(2.8)
        sec.right_margin = Cm(2.8)

    # ---- title page ----
    for _ in range(5):
        add_para(doc, "", indent=False)
    add_para(doc, "金钱心理学", size=30, bold=True, indent=False,
             align=WD_ALIGN_PARAGRAPH.CENTER, color=ACCENT, font=HEAD_FONT)
    add_para(doc, "The Psychology of Money", size=16, bold=False, indent=False,
             align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x66, 0x66, 0x66), font=HEAD_FONT)
    add_para(doc, "", indent=False)
    add_para(doc, "中 文 章 节 导 读", size=20, bold=True, indent=False,
             align=WD_ALIGN_PARAGRAPH.CENTER, font=HEAD_FONT)
    add_para(doc, "", indent=False)
    add_para(doc, "原书作者：Morgan Housel（摩根·豪泽尔）", size=13, indent=False,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "2026 年 9 月", size=12, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)

    # ---- reading notes ----
    doc.add_page_break()
    add_heading(doc, "导读说明", level=1)
    notes = [
        "本文档为《The Psychology of Money》（中译名《金钱心理学》，Morgan Housel 著，Harriman House 出版）的原创中文解读导读，供学习与研究参考。",
        "本文档并非原书的中文译本：内容为对各章思想与论点的概括提炼，以解读者的原创语言撰写，不替代阅读原书；如需完整内容，请阅读正版原书或其官方授权中译本。",
        "原书中的插图、图表等图片元素未收录于本文档，各章仅解读文字内容。",
        "财经专业名词与关键术语在首次出现时标注英文原文，例如：复利（Compounding）、安全边际（Margin of Safety）。",
        "文档结构：每章依次为「本章主旨」「核心观点」「关键故事与案例」「术语与概念」「实践启示」五个部分。",
    ]
    for n in notes:
        add_para(doc, n, size=12)
    add_para(doc, "全书结构：引言 + 20 章 + 后记。", size=12)

    # ---- parse digests ----
    n_head1 = 0
    for path in DIGESTS:
        with open(path, encoding="utf-8") as f:
            lines = f.read().splitlines()
        for raw in lines:
            line = raw.rstrip()
            if not line.strip():
                continue
            if line.startswith("# ") and not line.startswith("##"):
                add_heading(doc, line[2:].strip(), level=1)
                n_head1 += 1
            elif line.startswith("## "):
                add_heading(doc, line[3:].strip(), level=2)
            elif line.startswith("- "):
                add_bullet(doc, line[2:].strip())
            elif re.match(r"^\d+\.\s", line):
                add_bullet(doc, re.sub(r"^\d+\.\s+", "", line))
            else:
                add_para(doc, line.strip())

    doc.save(OUT)
    print("saved:", OUT, "| chapters:", n_head1)


if __name__ == "__main__":
    main()
